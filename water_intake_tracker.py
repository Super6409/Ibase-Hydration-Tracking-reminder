import tkinter as tk
from datetime import datetime
from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "water_intake.xlsx"
DAILY_GOAL = 2000
today = datetime.now().strftime("%Y-%m-%d")
total_water = 0

# Create the Excel file if needed
if not os.path.exists(FILE_NAME):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Water Intake"
    sheet.append(["Date", "Time", "Water (ml)"])
    workbook.save(FILE_NAME)

# Load today's total from Excel
workbook = load_workbook(FILE_NAME)
sheet = workbook.active
for row in sheet.iter_rows(min_row=2, values_only=True):
    if row[0] == today:
        total_water += row[2] or 0
workbook.close()

def add_water():
    global total_water

    try:
        amount = int(water_entry.get())
        if amount <= 0:
            message_label.config(text="Enter an amount above 0.")
            return
    except ValueError:
        message_label.config(text="Please enter a valid number.")
        return

    now = datetime.now()

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active
    sheet.append([today, now.strftime("%H:%M:%S"), amount])
    workbook.save(FILE_NAME)
    workbook.close()

    total_water += amount
    update_labels()
    message_label.config(text=f"Added {amount} ml.")
    water_entry.delete(0, tk.END)

def update_labels():
    total_label.config(text=f"Today's intake: {total_water} ml")
    remaining = max(0, DAILY_GOAL - total_water)
    goal_label.config(text=f"Goal: {DAILY_GOAL} ml | Remaining: {remaining} ml")

root = tk.Tk()
root.title("Water Intake Tracker")
root.geometry("330x230")
root.resizable(False, False)

tk.Label(root, text="Daily Water Intake", font=("Arial", 16, "bold")).pack(pady=12)

total_label = tk.Label(root, font=("Arial", 12))
total_label.pack(pady=4)

goal_label = tk.Label(root)
goal_label.pack(pady=4)

water_entry = tk.Entry(root, width=18, justify="center")
water_entry.pack(pady=8)
water_entry.insert(0, "Water in ml")

tk.Button(root, text="Add Water", command=add_water).pack(pady=4)

message_label = tk.Label(root, text="")
message_label.pack(pady=6)

update_labels()
root.mainloop()
